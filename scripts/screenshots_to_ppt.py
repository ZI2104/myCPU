"""
screenshots_to_ppt.py
将 slides_screenshots/ 下的 PNG 截图嵌入生成 4:3 PPT (10 x 7.5 inches)
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import os
import sys

SCREENSHOTS_DIR = r"d:\code\myCPU\docs\presentations\slides_screenshots"
OUTPUT_FILE     = r"d:\code\myCPU\docs\presentations\myCPU_演示_4x3.pptx"

SLIDE_TITLES = [
    "封面 - myCPU 项目开题汇报",
    "技术选型",
    "系统架构总览",
    "六级流水线设计",
    "特权级与中断系统",
    "四大技术亮点",
    "已完成工作",
    "项目进度规划",
    "扩展亮点",
    "总结",
]

def main():
    # 4:3 layout: 10 x 7.5 inches
    prs = Presentation()
    prs.slide_width  = Inches(10)
    prs.slide_height = Inches(7.5)

    # Get sorted screenshots
    files = sorted([
        f for f in os.listdir(SCREENSHOTS_DIR)
        if f.lower().endswith('.png')
    ])
    print(f"Found {len(files)} screenshots")

    blank_layout = prs.slide_layouts[6]  # completely blank layout

    for i, fname in enumerate(files):
        img_path = os.path.join(SCREENSHOTS_DIR, fname)
        slide = prs.slides.add_slide(blank_layout)

        # Fill background with dark color matching the HTML theme
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(0x0A, 0x0E, 0x1A)

        # Add image covering the full slide
        slide.shapes.add_picture(
            img_path,
            left   = Inches(0),
            top    = Inches(0),
            width  = Inches(10),
            height = Inches(7.5)
        )

        title = SLIDE_TITLES[i] if i < len(SLIDE_TITLES) else fname
        print(f"  ✓ Slide {i+1}: {title}")

    prs.save(OUTPUT_FILE)
    print(f"\nPPT saved to: {OUTPUT_FILE}")
    size_kb = os.path.getsize(OUTPUT_FILE) / 1024
    print(f"File size: {size_kb:.1f} KB")

if __name__ == "__main__":
    main()
